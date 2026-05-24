#!/usr/bin/env python3
"""
Analyze XLSM file structure and formatting
"""

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
import json

def analyze_xlsm(filename):
    """Analyze XLSM file structure."""
    wb = load_workbook(filename=filename, read_only=False, data_only=False)
    
    print("="*80)
    print(f"XLSM Analysis: {filename}")
    print("="*80)
    
    # Sheet names
    print(f"\n📊 Sheet Names: {wb.sheetnames}")
    
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f"\n{'='*80}")
        print(f"Sheet: {sheet_name}")
        print(f"{'='*80}")
        print(f"Dimensions: {ws.dimensions}")
        print(f"Max Row: {ws.max_row}, Max Column: {ws.max_column}")
        
        # Analyze headers
        print("\n📋 Headers (Row 1):")
        headers = []
        for cell in ws[1]:
            if cell.value:
                headers.append(cell.value)
                print(f"  • {cell.coordinate}: {cell.value}")
                print(f"    - Font: {cell.font.name}, Size: {cell.font.size}, Bold: {cell.font.bold}")
                print(f"    - Font Color: {cell.font.color.rgb if cell.font.color else 'None'}")
                print(f"    - Fill Color: {cell.fill.fgColor.rgb if cell.fill.fgColor else 'None'}")
                print(f"    - Alignment: H={cell.alignment.horizontal}, V={cell.alignment.vertical}")
        
        # Sample data rows
        print("\n📊 Sample Data (First 5 rows):")
        for row_idx in range(2, min(7, ws.max_row + 1)):
            row_data = []
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                row_data.append(cell.value)
            print(f"  Row {row_idx}: {row_data}")
        
        # Data types analysis
        print("\n🔍 Data Type Analysis:")
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=2, column=col_idx)
            print(f"  • {header}: {type(cell.value).__name__}")

if __name__ == "__main__":
    analyze_xlsm("d868m.xlsm")
